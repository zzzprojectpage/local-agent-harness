import tempfile
import os
import subprocess
import unittest
import hashlib
import zipfile
from pathlib import Path

from local_harness.workspace import Workspace, WorkspaceError


class WorkspaceTests(unittest.TestCase):
    def test_cell_writes_preserve_text_identifiers_and_do_not_auto_convert_dates(self):
        with tempfile.TemporaryDirectory() as temp:
            workspace = Workspace(temp, autonomous=True)
            workspace.create_workbook("codes.xlsx", "Data")
            workspace.set_cells("codes.xlsx", "Data", "A1:C1", [["00123", "2026-09-30", "+42"]])
            self.assertEqual(workspace.read_excel("codes.xlsx", "Data", "A1:C1")["values"], [["00123", "2026-09-30", "+42"]])

    def test_failed_edit_keeps_the_published_working_copy_and_original_unchanged(self):
        import openpyxl
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            book = openpyxl.Workbook()
            book.active.title = "Data"
            book.active["A1"] = 1
            book.save(root / "data.xlsx")
            book.close()
            original = (root / "data.xlsx").read_bytes()
            workspace = Workspace(root, autonomous=True)
            result = workspace.set_cells("data.xlsx", "Data", "A1", [[2]])
            output = root / result["output_path"]
            previous = output.read_bytes()
            with self.assertRaises(ValueError):
                workspace.set_cells("data.xlsx", "Missing", "A1", [[3]])
            self.assertEqual(output.read_bytes(), previous)
            self.assertEqual((root / "data.xlsx").read_bytes(), original)
            self.assertEqual(list(root.glob(".harness-stage-*")), [])

    def test_autonomous_text_edits_need_no_callback_and_preserve_originals(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            (root / "notes.txt").write_text("original", encoding="utf-8")
            workspace = Workspace(root, autonomous=True)
            result = workspace.write_file("notes.txt", "first")
            self.assertEqual(result["status"], "written")
            self.assertNotEqual(result["output_path"], "notes.txt")
            second = workspace.write_file("notes.txt", "second")
            self.assertEqual(second["output_path"], result["output_path"])
            self.assertEqual(workspace.read_file("notes.txt"), "1: second")
            self.assertEqual((root / "notes.txt").read_text(encoding="utf-8"), "original")

    def test_edits_binary_workbooks_in_their_original_format(self):
        import shutil
        for suffix in ("xls", "xlsb"):
            with self.subTest(format=suffix), tempfile.TemporaryDirectory() as temp:
                root = Path(temp)
                target = root / ("sales." + suffix)
                shutil.copyfile(Path(__file__).parent / "fixtures" / ("sample." + suffix), target)
                original = target.read_bytes()
                workspace = Workspace(root, autonomous=True)
                result = workspace.set_cells(target.name, "Sales", "B2", [[11]])
                self.assertTrue(result["output_path"].endswith("." + suffix))
                self.assertEqual(workspace.read_excel(target.name, "Sales", "B2")["values"], [[11]])
                self.assertEqual(target.read_bytes(), original)

    def test_ooxml_templates_and_macro_workbooks_can_be_read(self):
        import openpyxl
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            for suffix in ("xlsx", "xlsm", "xltx", "xltm", "xlam"):
                with self.subTest(format=suffix):
                    book = openpyxl.Workbook()
                    book.template = suffix in ("xltx", "xltm")
                    book.active.title = "Data"
                    book.active["A1"] = "Template fact"
                    path = root / ("sample." + suffix)
                    book.save(path)
                    book.close()
                    self.assertEqual(Workspace(root).read_excel(path.name, "Data", "A1")["values"], [["Template fact"]])

    def test_excel_permissions_ranges_and_formula_guards_are_enforced(self):
        import openpyxl
        import threading
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            book = openpyxl.Workbook()
            book.save(root / "book.xlsx")
            book.close()
            original = (root / "book.xlsx").read_bytes()
            with self.assertRaises(WorkspaceError):
                Workspace(root).set_cells("book.xlsx", cell_range="A1", values=[[1]])
            workspace = Workspace(root, autonomous=True)
            for cell_range in ("A0", "XFE1", "A1:XFD1048576", "A2:A1", "Sheet1!A1"):
                with self.subTest(range=cell_range), self.assertRaises(ValueError):
                    workspace.read_excel("book.xlsx", cell_range=cell_range)
            with self.assertRaises(ValueError):
                workspace.set_cells("book.xlsx", cell_range="A1:B1", values=[[1]])
            with self.assertRaises(ValueError):
                workspace.set_cells("book.xlsx", cell_range="A1", values=[['=WEBSERVICE("https://example.invalid")']])
            stopped = threading.Event()
            stopped.set()
            with self.assertRaises(ValueError):
                workspace.set_cells("book.xlsx", cell_range="A1", values=[[1]], cancelled=stopped)
            with self.assertRaises(WorkspaceError):
                workspace.run_vba_macro("book.xlsx", "Existing", "Run")
            with self.assertRaises(WorkspaceError):
                workspace.read_excel("../outside.xlsx")
            self.assertEqual((root / "book.xlsx").read_bytes(), original)

    def test_fill_formulas_and_formatting_are_saved_with_relative_references(self):
        import openpyxl
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            book = openpyxl.Workbook()
            book.active.title = "Data"
            book.active.append([2, 3])
            book.active.append([4, 5])
            book.save(root / "book.xlsx")
            book.close()
            workspace = Workspace(root, autonomous=True)
            workspace.fill_formula("book.xlsx", "Data", "C1:C2", "=A1*B1")
            result = workspace.format_cells("book.xlsx", "Data", "C1:C2", number_format="0.00", bold=True, fill_color="#D1FAE5")
            read = workspace.read_excel("book.xlsx", "Data", "C1:C2")
            self.assertEqual(read["values"], [[6], [20]])
            self.assertEqual(read["formulas"], {"C1": "=A1*B1", "C2": "=A2*B2"})
            reopened = openpyxl.load_workbook(root / result["output_path"])
            try:
                self.assertTrue(reopened["Data"]["C2"].font.bold)
                self.assertEqual(reopened["Data"]["C2"].number_format, "0.00")
            finally:
                reopened.close()

    def test_agent_can_create_a_workbook_and_add_a_sheet_without_overwriting(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            workspace = Workspace(root, autonomous=True)
            created = workspace.create_workbook("report.xlsx", sheet="Data")
            self.assertEqual(created["output_path"], "report.xlsx")
            workspace.add_sheet("report.xlsx", "Summary")
            self.assertEqual([s["name"] for s in workspace.inspect_workbook("report.xlsx")["sheets"]], ["Data", "Summary"])
            with self.assertRaises((ValueError, WorkspaceError)):
                workspace.create_workbook("report.xlsx")
            with self.assertRaises((ValueError, WorkspaceError)):
                workspace.add_sheet("report.xlsx", "Summary")

    def test_creates_and_runs_vba_on_an_xlsm_copy_without_confirmation(self):
        import openpyxl
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            book = openpyxl.Workbook()
            book.active.title = "Data"
            book.save(root / "macro.xlsx")
            book.close()
            original = (root / "macro.xlsx").read_bytes()
            workspace = Workspace(root, autonomous=True, allow_vba=True)
            code = 'Option Explicit\nPublic Sub Auto_Open()\n ThisWorkbook.Worksheets("Data").Range("D1").Value = 666\nEnd Sub\nPublic Sub SetResult()\n ThisWorkbook.Worksheets("Data").Range("B2").Value = 42\nEnd Sub'
            result = workspace.create_vba_module("macro.xlsx", "Results", code)
            self.assertTrue(result["output_path"].endswith(".xlsm"))
            self.assertTrue(result["vba_created"])
            # Execution provenance must survive use of the returned output path.
            result = workspace.run_vba_macro(result["output_path"], "Results", "SetResult")
            self.assertTrue(result["macros_executed"])
            self.assertEqual(workspace.read_excel("macro.xlsx", "Data", "B2")["values"], [[42]])
            self.assertEqual(workspace.read_excel("macro.xlsx", "Data", "D1")["values"], [[None]], "Opening and explicitly running SetResult must not invoke Auto_Open.")
            script = root / "Imported.bas"
            script.write_text('Attribute VB_Name = "Imported"\nOption Explicit\nPublic Sub SetOther()\n ThisWorkbook.Worksheets("Data").Range("C3").Value = 73\nEnd Sub', encoding="utf-8")
            imported = workspace.import_vba_module("macro.xlsx", "Imported.bas")
            self.assertEqual(imported["module"], "Imported")
            workspace.run_vba_macro("macro.xlsx", "Imported", "SetOther")
            self.assertEqual(workspace.read_excel("macro.xlsx", "Data", "C3")["values"], [[73]])
            self.assertEqual((root / "macro.xlsx").read_bytes(), original)

    def test_reads_real_binary_excel_formats_without_modifying_them(self):
        import shutil
        for suffix in ("xls", "xlsb"):
            with self.subTest(format=suffix), tempfile.TemporaryDirectory() as temp:
                root = Path(temp)
                target = root / ("sales." + suffix)
                shutil.copyfile(Path(__file__).parent / "fixtures" / ("sample." + suffix), target)
                original = target.read_bytes()
                workspace = Workspace(root)
                self.assertEqual(workspace.inspect_workbook(target.name)["sheets"][0]["name"], "Sales")
                result = workspace.read_excel(target.name, "Sales", "A1:C2")
                self.assertEqual(result["values"], [["Item", "Quantity", "Price"], ["Apples", 3, 7]])
                self.assertEqual(target.read_bytes(), original)

    def test_autonomous_excel_values_and_formulas_use_a_working_copy(self):
        import openpyxl
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            original = root / "sales.xlsx"
            book = openpyxl.Workbook()
            book.active.title = "Sales"
            book.active.append(["Item", "Quantity", "Price"])
            book.active.append(["Apples", 3, 7])
            book.save(original)
            book.close()
            digest = hashlib.sha256(original.read_bytes()).hexdigest()
            workspace = Workspace(root, autonomous=True)
            changed = workspace.set_cells("sales.xlsx", sheet="Sales", cell_range="B2", values=[[5]])
            self.assertEqual(changed["status"], "written")
            self.assertTrue(changed["saved_read_verified"])
            self.assertEqual(changed["saved_preview"]["values"], [[5]])
            self.assertNotEqual(changed["output_path"], "sales.xlsx")
            workspace.set_cells("sales.xlsx", sheet="Sales", cell_range="D2", values=[["=B2*C2"]])
            result = workspace.read_excel("sales.xlsx", sheet="Sales", cell_range="A1:D2")
            self.assertEqual(result["values"][1][1], 5)
            self.assertEqual(result["values"][1][3], 35)
            self.assertEqual(result["formulas"]["D2"], "=B2*C2")
            self.assertEqual(hashlib.sha256(original.read_bytes()).hexdigest(), digest)

    def test_native_pivot_table_is_created_without_a_confirmation_callback(self):
        import openpyxl
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            book = openpyxl.Workbook()
            book.active.title = "Data"
            for row in (("Category", "Amount"), ("A", 10), ("A", 20), ("B", 5)):
                book.active.append(row)
            book.save(root / "data.xlsx")
            book.close()
            workspace = Workspace(root, autonomous=True)
            result = workspace.create_pivot("data.xlsx", source_sheet="Data", source_range="A1:B4",
                                            row_fields=["Category"], data_fields=[{"field": "Amount", "aggregation": "sum"}],
                                            target_sheet="Summary", name="CategoryTotals")
            self.assertEqual(result["status"], "written")
            output = root / result["output_path"]
            with zipfile.ZipFile(output) as archive:
                self.assertIn("xl/pivotTables/pivotTable1.xml", archive.namelist())
            values = workspace.read_excel("data.xlsx", "Summary", "A3:B7")["values"]
            self.assertIn(["A", 30], values)
            self.assertIn(["B", 5], values)
            self.assertIn(["Grand Total", 35], values)

    def test_reads_xlsx_cells_and_formulas_without_changing_the_workbook(self):
        import openpyxl
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            path = root / "sales.xlsx"
            book = openpyxl.Workbook()
            book.active.title = "Sales"
            book.active.append(["Item", "Quantity", "Price"])
            book.active.append(["Apples", 3, 7])
            book.active["D2"] = "=B2*C2"
            book.save(path)
            book.close()
            original = hashlib.sha256(path.read_bytes()).hexdigest()
            workspace = Workspace(root)
            self.assertEqual(workspace.list_files(), ["sales.xlsx"])
            metadata = workspace.inspect_workbook("sales.xlsx")
            self.assertEqual(metadata["sheets"][0]["name"], "Sales")
            result = workspace.read_excel("sales.xlsx", sheet="Sales", cell_range="A1:D2")
            self.assertEqual(result["values"][1][:3], ["Apples", 3, 7])
            self.assertEqual(result["formulas"]["D2"], "=B2*C2")
            self.assertIn("Sales", workspace.read_file("sales.xlsx"))
            self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(), original)

    def test_read_is_numbered_and_cannot_escape_the_selected_folder(self):
        with tempfile.TemporaryDirectory() as temp:
            parent = Path(temp)
            root = parent / "project"
            root.mkdir()
            (root / "notes.txt").write_text("apples\npears\n", encoding="utf-8")
            (parent / "private.txt").write_text("outside", encoding="utf-8")
            workspace = Workspace(root)
            self.assertEqual(workspace.read_file("notes.txt"), "1: apples\n2: pears")
            with self.assertRaises(WorkspaceError):
                workspace.read_file("../private.txt")
            with self.assertRaises(WorkspaceError):
                workspace.read_file(str(parent / "private.txt"))

    def test_search_skips_secret_and_binary_files_and_reads_unicode(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            (root / "report.csv").write_text("name,total\nDiana,42\n", encoding="utf-8")
            (root / ".env").write_text("name=hidden", encoding="utf-8")
            (root / "binary.txt").write_bytes(b"name\x00hidden")
            (root / "notes.txt").write_text("Bună ziua\n", encoding="utf-16")
            workspace = Workspace(root)
            self.assertEqual(workspace.list_files(), ["binary.txt", "notes.txt", "report.csv"])
            self.assertEqual(workspace.search_text("Diana"), [
                {"path": "report.csv", "line": 2, "text": "Diana,42"}
            ])
            self.assertEqual(workspace.read_file("notes.txt"), "1: Bună ziua")
            for path in (".env", "binary.txt"):
                with self.assertRaises(WorkspaceError):
                    workspace.read_file(path)

    def test_writes_require_approval_and_preserve_an_existing_file_in_backup(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            target = root / "notes.txt"
            target.write_text("old\n", encoding="utf-8")
            workspace = Workspace(root)
            denied = workspace.write_file("notes.txt", "new\n", approve=lambda request: False)
            self.assertEqual(denied["status"], "denied")
            self.assertEqual(target.read_text(encoding="utf-8"), "old\n")
            requests = []

            def approve(request):
                requests.append(request)
                return True

            result = workspace.write_file("notes.txt", "new\n", approve=approve)
            self.assertEqual(result["status"], "written")
            self.assertEqual(target.read_text(encoding="utf-8"), "new\n")
            self.assertEqual((root / result["backup"]).read_text(encoding="utf-8"), "old\n")
            self.assertIn("-old", requests[0].diff)
            self.assertNotIn(result["backup"], workspace.list_files())
            with self.assertRaises(WorkspaceError):
                workspace.write_file("book.xlsx", "corrupt", approve=approve)

    def test_file_changed_during_approval_is_not_overwritten(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            target = root / "notes.txt"
            target.write_text("original", encoding="utf-8")

            def approve(request):
                target.write_text("human change", encoding="utf-8")
                return True

            with self.assertRaises(WorkspaceError):
                Workspace(root).write_file("notes.txt", "model change", approve=approve)
            self.assertEqual(target.read_text(encoding="utf-8"), "human change")

    def test_external_directory_links_are_blocked(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp) / "project"
            outside = Path(temp) / "outside"
            root.mkdir()
            outside.mkdir()
            (outside / "notes.txt").write_text("external", encoding="utf-8")
            try:
                (root / "link").symlink_to(outside, target_is_directory=True)
            except OSError:
                if os.name != "nt":
                    self.skipTest("Symlinks unavailable on this machine.")
                result = subprocess.run(["cmd", "/c", "mklink", "/J", str(root / "link"), str(outside)], capture_output=True)
                if result.returncode:
                    self.skipTest("Neither symlinks nor Windows directory junctions are available.")
            with self.assertRaises(WorkspaceError):
                Workspace(root).read_file("link/notes.txt")
            self.assertEqual(Workspace(root).list_files(), [])
            with self.assertRaises(WorkspaceError):
                Workspace(root).write_file("link/new.txt", "outside", approve=lambda request: True)


if __name__ == "__main__":
    unittest.main()
